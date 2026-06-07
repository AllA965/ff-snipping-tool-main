"""
Office导出模块
支持导出到Word、Excel、PowerPoint，并自动打开
"""
import os
import sys
import tempfile
import subprocess
from pathlib import Path
from typing import Optional, Tuple
from PySide6.QtGui import QPixmap, QImage
from PySide6.QtCore import QBuffer, QIODevice


class OfficeExporter:
    """Office导出器"""
    
    @staticmethod
    def get_office_paths() -> dict:
        """获取Office程序路径"""
        paths = {
            'word': None,
            'excel': None,
            'powerpoint': None,
            'paint': None
        }
        
        if sys.platform == 'win32':
            # Windows Office路径
            office_locations = [
                r"C:\Program Files\Microsoft Office\root\Office16",
                r"C:\Program Files (x86)\Microsoft Office\root\Office16",
                r"C:\Program Files\Microsoft Office\Office16",
                r"C:\Program Files (x86)\Microsoft Office\Office16",
                r"C:\Program Files\Microsoft Office\Office15",
                r"C:\Program Files (x86)\Microsoft Office\Office15",
            ]
            
            for loc in office_locations:
                if os.path.exists(loc):
                    word_path = os.path.join(loc, "WINWORD.EXE")
                    excel_path = os.path.join(loc, "EXCEL.EXE")
                    ppt_path = os.path.join(loc, "POWERPNT.EXE")
                    
                    if os.path.exists(word_path):
                        paths['word'] = word_path
                    if os.path.exists(excel_path):
                        paths['excel'] = excel_path
                    if os.path.exists(ppt_path):
                        paths['powerpoint'] = ppt_path
                    break
            
            # Paint路径
            paint_locations = [
                r"C:\Windows\System32\mspaint.exe",
                os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\WindowsApps\mspaint.exe"),
                r"C:\Windows\SysWOW64\mspaint.exe",
            ]
            for paint_path in paint_locations:
                if os.path.exists(paint_path):
                    paths['paint'] = paint_path
                    break
        
        elif sys.platform == 'darwin':
            # macOS Office路径
            if os.path.exists("/Applications/Microsoft Word.app"):
                paths['word'] = "/Applications/Microsoft Word.app"
            if os.path.exists("/Applications/Microsoft Excel.app"):
                paths['excel'] = "/Applications/Microsoft Excel.app"
            if os.path.exists("/Applications/Microsoft PowerPoint.app"):
                paths['powerpoint'] = "/Applications/Microsoft PowerPoint.app"
        
        return paths
    
    @staticmethod
    def is_office_available() -> dict:
        """检查Office组件是否可用"""
        paths = OfficeExporter.get_office_paths()
        return {
            'word': paths['word'] is not None,
            'excel': paths['excel'] is not None,
            'powerpoint': paths['powerpoint'] is not None,
            'paint': paths['paint'] is not None
        }
    
    @staticmethod
    def pixmap_to_bytes(pixmap: QPixmap, format: str = "PNG") -> bytes:
        """将QPixmap转换为字节数据"""
        buffer = QBuffer()
        buffer.open(QIODevice.OpenModeFlag.WriteOnly)
        pixmap.save(buffer, format)
        return bytes(buffer.data())
    
    @staticmethod
    def export_to_word(pixmap: QPixmap, filename: str = None) -> Tuple[bool, str]:
        """导出到Word文档"""
        try:
            from docx import Document
            from docx.shared import Inches
            from io import BytesIO
            
            # 创建文档
            doc = Document()
            
            # 添加标题
            doc.add_heading('截图内容', 0)
            
            # 将图片转换为字节流
            img_bytes = OfficeExporter.pixmap_to_bytes(pixmap)
            img_stream = BytesIO(img_bytes)
            
            # 计算合适的宽度（最大6英寸）
            width = pixmap.width()
            height = pixmap.height()
            max_width = 6.0
            
            if width > 0:
                aspect = height / width
                doc_width = min(width / 96, max_width)  # 96 DPI
                doc.add_picture(img_stream, width=Inches(doc_width))
            
            # 保存文件
            if not filename:
                filename = tempfile.mktemp(suffix='.docx')
            
            doc.save(filename)
            
            # 打开Word
            OfficeExporter._open_with_office('word', filename)
            
            return True, filename
            
        except ImportError:
            return False, "需要安装 python-docx 库: pip install python-docx"
        except Exception as e:
            return False, f"导出失败: {str(e)}"
    
    @staticmethod
    def export_to_excel(pixmap: QPixmap, filename: str = None) -> Tuple[bool, str]:
        """导出到Excel文档"""
        try:
            from openpyxl import Workbook
            from openpyxl.drawing.image import Image as XLImage
            from io import BytesIO
            
            # 创建工作簿
            wb = Workbook()
            ws = wb.active
            ws.title = "截图"
            
            # 保存临时图片
            img_bytes = OfficeExporter.pixmap_to_bytes(pixmap)
            img_stream = BytesIO(img_bytes)
            
            # 添加图片
            img = XLImage(img_stream)
            
            # 调整大小
            max_width = 800
            if img.width > max_width:
                ratio = max_width / img.width
                img.width = int(img.width * ratio)
                img.height = int(img.height * ratio)
            
            ws.add_image(img, 'A1')
            
            # 保存文件
            if not filename:
                filename = tempfile.mktemp(suffix='.xlsx')
            
            wb.save(filename)
            
            # 打开Excel
            OfficeExporter._open_with_office('excel', filename)
            
            return True, filename
            
        except ImportError:
            return False, "需要安装 openpyxl 库: pip install openpyxl"
        except Exception as e:
            return False, f"导出失败: {str(e)}"
    
    @staticmethod
    def export_to_powerpoint(pixmap: QPixmap, filename: str = None) -> Tuple[bool, str]:
        """导出到PowerPoint文档"""
        try:
            from pptx import Presentation
            from pptx.util import Inches
            from io import BytesIO
            
            # 创建演示文稿
            prs = Presentation()
            
            # 添加空白幻灯片
            blank_layout = prs.slide_layouts[6]  # 空白布局
            slide = prs.slides.add_slide(blank_layout)
            
            # 将图片转换为字节流
            img_bytes = OfficeExporter.pixmap_to_bytes(pixmap)
            img_stream = BytesIO(img_bytes)
            
            # 计算图片位置和大小（居中显示）
            slide_width = prs.slide_width.inches
            slide_height = prs.slide_height.inches
            
            img_width = pixmap.width() / 96  # 转换为英寸
            img_height = pixmap.height() / 96
            
            # 缩放以适应幻灯片
            scale = min(slide_width / img_width, slide_height / img_height, 1.0) * 0.9
            final_width = img_width * scale
            final_height = img_height * scale
            
            # 居中
            left = (slide_width - final_width) / 2
            top = (slide_height - final_height) / 2
            
            slide.shapes.add_picture(
                img_stream, 
                Inches(left), 
                Inches(top),
                Inches(final_width),
                Inches(final_height)
            )
            
            # 保存文件
            if not filename:
                filename = tempfile.mktemp(suffix='.pptx')
            
            prs.save(filename)
            
            # 打开PowerPoint
            OfficeExporter._open_with_office('powerpoint', filename)
            
            return True, filename
            
        except ImportError:
            return False, "需要安装 python-pptx 库: pip install python-pptx"
        except Exception as e:
            return False, f"导出失败: {str(e)}"
    
    @staticmethod
    def send_to_paint(pixmap: QPixmap) -> Tuple[bool, str]:
        """发送到画图程序"""
        try:
            # 保存临时文件
            temp_file = tempfile.mktemp(suffix='.png')
            pixmap.save(temp_file, "PNG")
            
            # 打开画图
            OfficeExporter._open_with_office('paint', temp_file)
            
            return True, temp_file
            
        except Exception as e:
            return False, f"发送失败: {str(e)}"
    
    @staticmethod
    def send_to_photoshop(pixmap: QPixmap) -> Tuple[bool, str]:
        """发送到Photoshop"""
        try:
            # 保存为PSD兼容的PNG（保持透明度）
            temp_file = tempfile.mktemp(suffix='.png')
            pixmap.save(temp_file, "PNG")
            
            # 查找Photoshop
            ps_path = None
            if sys.platform == 'win32':
                ps_locations = [
                    r"C:\Program Files\Adobe\Adobe Photoshop 2024\Photoshop.exe",
                    r"C:\Program Files\Adobe\Adobe Photoshop 2023\Photoshop.exe",
                    r"C:\Program Files\Adobe\Adobe Photoshop CC 2022\Photoshop.exe",
                    r"C:\Program Files\Adobe\Adobe Photoshop CC 2021\Photoshop.exe",
                    r"C:\Program Files\Adobe\Adobe Photoshop CC 2020\Photoshop.exe",
                ]
                for loc in ps_locations:
                    if os.path.exists(loc):
                        ps_path = loc
                        break
            elif sys.platform == 'darwin':
                ps_path = "/Applications/Adobe Photoshop 2024/Adobe Photoshop 2024.app"
            
            if ps_path:
                if sys.platform == 'win32':
                    subprocess.Popen([ps_path, temp_file])
                else:
                    subprocess.Popen(['open', '-a', ps_path, temp_file])
                return True, temp_file
            else:
                # 尝试使用系统关联打开
                if sys.platform == 'win32':
                    os.startfile(temp_file)
                else:
                    subprocess.Popen(['open', temp_file])
                return True, temp_file
                
        except Exception as e:
            return False, f"发送失败: {str(e)}"
    
    @staticmethod
    def send_to_gimp(pixmap: QPixmap) -> Tuple[bool, str]:
        """发送到GIMP"""
        try:
            temp_file = tempfile.mktemp(suffix='.png')
            pixmap.save(temp_file, "PNG")
            
            gimp_path = None
            if sys.platform == 'win32':
                gimp_locations = [
                    r"C:\Program Files\GIMP 2\bin\gimp-2.10.exe",
                    r"C:\Program Files (x86)\GIMP 2\bin\gimp-2.10.exe",
                ]
                for loc in gimp_locations:
                    if os.path.exists(loc):
                        gimp_path = loc
                        break
            elif sys.platform == 'darwin':
                gimp_path = "/Applications/GIMP-2.10.app"
            
            if gimp_path:
                if sys.platform == 'win32':
                    subprocess.Popen([gimp_path, temp_file])
                else:
                    subprocess.Popen(['open', '-a', gimp_path, temp_file])
                return True, temp_file
            else:
                return False, "未找到GIMP程序"
                
        except Exception as e:
            return False, f"发送失败: {str(e)}"
    
    @staticmethod
    def _open_with_office(app: str, filename: str):
        """使用Office程序打开文件"""
        paths = OfficeExporter.get_office_paths()
        app_path = paths.get(app)
        
        if app_path:
            if sys.platform == 'win32':
                subprocess.Popen([app_path, filename])
            elif sys.platform == 'darwin':
                subprocess.Popen(['open', '-a', app_path, filename])
        else:
            # 使用系统默认程序打开
            if sys.platform == 'win32':
                os.startfile(filename)
            elif sys.platform == 'darwin':
                subprocess.Popen(['open', filename])
            else:
                subprocess.Popen(['xdg-open', filename])
